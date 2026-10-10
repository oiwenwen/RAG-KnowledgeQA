import os
from pathlib import Path
import gradio as gr

# ========== 配置区 ==========
os.environ["DASHSCOPE_API_KEY"] = "sk-ws-H.PRXMYEM.alfm.MEUCIAKq0dD38mAqNJHdZMP4x3odPurVl8ilihlk3PmYKcCzAiEA9jiPGeOQo6Axy6XoZTLnOtcdWbCLylU2xTJDDrYNf6g"

# 向量库临时保存路径
VECTOR_STORE_PATH = "temp_faiss_index"
# 保留的历史对话轮数
MAX_HISTORY_ROUND = 3

# ========== 导入依赖 ==========
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import DashScopeEmbeddings
from pypdf import PdfReader
from rank_bm25 import BM25Okapi
from langchain_core.documents import Document
import docx
import dashscope
import pickle

# ========== 工具函数 ==========
def load_document(file_path):
    """读取不同格式文档"""
    file_path = Path(file_path)
    suffix = file_path.suffix.lower()
    
    if suffix == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    
    elif suffix == ".pdf":
        text = ""
        reader = PdfReader(file_path)
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    
    elif suffix == ".docx":
        doc = docx.Document(file_path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        return text
    
    else:
        raise ValueError(f"不支持的文件格式：{suffix}")

def build_vector_store(file_obj):
    """根据上传的文档构建向量库，同时清空对话历史"""
    # 清理旧的向量库
    if os.path.exists(VECTOR_STORE_PATH):
        import shutil
        shutil.rmtree(VECTOR_STORE_PATH)
    if os.path.exists("chunks.pkl"):
        os.remove("chunks.pkl")
    
        # 初始化总分块列表
    all_chunks = []
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=50,
        separators=["\n\n", "\n", "。", "，", " "]
    )
    
    # 循环处理每个上传的文档
    for file_item in file_obj:
        raw_text = load_document(file_item.name)
        chunks = text_splitter.split_text(raw_text)
        all_chunks.extend(chunks)
    
    chunks = all_chunks

    
    # 构建向量库
    embeddings = DashScopeEmbeddings(model="text-embedding-v2")
    vector_store = FAISS.from_texts(chunks, embeddings)
    vector_store.save_local(VECTOR_STORE_PATH)
    
    # 保存文本块副本，供BM25关键词检索使用
    with open("chunks.pkl", "wb") as f:
        pickle.dump(chunks, f)
    
    # 返回：上传状态 + 空对话历史（标准字典格式）
    return f"✅ 知识库构建完成！共处理 {len(file_obj)} 个文档，解析 {len(chunks)} 个文本块。", []


def answer_question(query, chat_history):
    # ========== 第1步：先显示用户问题 ==========
    chat_history.append({"role": "user", "content": query})
    yield chat_history, "", "", ""

    # ========== 第2步：检查向量库 ==========
    if not os.path.exists(VECTOR_STORE_PATH):
        chat_history.append({"role": "assistant", "content": "⚠️ 请先上传文档！"})
        yield chat_history, "", "", ""
        return

    # ========== 第3步：混合检索上下文 ==========
    # 加载向量库
    embeddings = DashScopeEmbeddings(model="text-embedding-v2")
    vector_store = FAISS.load_local(
        VECTOR_STORE_PATH, 
        embeddings,
        allow_dangerous_deserialization=True
    )

    # 1. 向量语义检索
    vector_docs = vector_store.similarity_search(query, k=3)
    
    # 2. BM25关键词检索
    with open("chunks.pkl", "rb") as f:
        all_chunks = pickle.load(f)
    tokenized_corpus = [list(doc) for doc in all_chunks]
    bm25 = BM25Okapi(tokenized_corpus)
    tokenized_query = list(query)
    bm25_scores = bm25.get_scores(tokenized_query)
    top_bm25_idx = sorted(range(len(bm25_scores)), key=lambda i: bm25_scores[i], reverse=True)[:3]
    bm25_docs = [Document(page_content=all_chunks[i]) for i in top_bm25_idx]
    
    # 3. 合并去重
    combined_docs = []
    seen_contents = set()
    for doc in vector_docs + bm25_docs:
        if doc.page_content not in seen_contents:
            seen_contents.add(doc.page_content)
            combined_docs.append(doc)
    related_docs = combined_docs

    # ========== 第4步：构建提示词 ==========
    # 给上下文段落编号
    context = ""
    for i, doc in enumerate(related_docs, 1):
        context += f"[{i}] {doc.page_content}\n\n"
    
    # 拼接历史对话
    history_text = ""
    history_messages = chat_history[-MAX_HISTORY_ROUND * 2:]
    for msg in history_messages:
        if msg["role"] == "user":
            history_text += f"用户：{msg['content']}\n"
        else:
            history_text += f"助手：{msg['content']}\n"
    
    # 最终提示词
    prompt = f"""请你结合对话历史，仅根据下面的上下文内容回答用户的问题。
规则：
1. 只能使用上下文里的信息，禁止编造上下文没有的内容
2. 每个知识点回答结束后，必须用 [编号] 标注对应的来源段落，例如：电子信息学院开设计算机专业[1]
3. 如果多个段落都提到同一个点，可以标注多个编号，例如：学校有两个校区[1][3]
4. 如果上下文里没有相关信息，就回答"知识库中暂无相关内容"
5. 结合对话历史理解用户的问题，保持对话连贯性

对话历史：
{history_text}

上下文：
{context}

用户当前问题：{query}

请给出简洁准确的回答："""

    # ========== 第5步：流式调用大模型 ==========
    responses = dashscope.Generation.call(
        model="qwen-turbo",
        prompt=prompt,
        result_format="text",
        stream=True
    )

    answer = ""
    usage = None
    # 用副本操作，不污染原始对话历史
    temp_chat = chat_history.copy()
    # 预先占位一条空的助手消息
    temp_chat.append({"role": "assistant", "content": ""})

    for resp in responses:
        if resp.output is None:
            temp_chat[-1]["content"] = f"❌ 调用失败：{resp.message}"
            yield temp_chat, "", "", ""
            return
        
         # 直接使用最新全量文本，不要累加
        if resp.output.text:
            answer = resp.output.text

        
        # 记录Token统计
        if resp.usage:
            usage = resp.usage

        # 实时更新助手消息
        temp_chat[-1]["content"] = answer
        yield temp_chat, "", "", ""

    # ========== 第6步：生成完成，返回完整结果 ==========
    # 同步最终对话
    chat_history[:] = temp_chat

    # 格式化引用原文
    ref_text = "📄 引用原文（与回答中编号对应）：\n\n"
    for i, doc in enumerate(related_docs, 1):
        ref_text += f"[{i}] {doc.page_content.strip()}\n\n"
    
    # 格式化Token统计
    token_text = f"""📊 Token消耗统计
输入Token：{usage.input_tokens}
输出Token：{usage.output_tokens}
本次总计：{usage.total_tokens} Tokens"""

    # 最终返回完整结果
    yield chat_history, "", ref_text, token_text

def clear_knowledge_base():
    """清空整个知识库：删除索引文件+清空对话+重置状态"""
    # 删除向量库文件夹
    if os.path.exists(VECTOR_STORE_PATH):
        import shutil
        shutil.rmtree(VECTOR_STORE_PATH)
    # 删除文本块缓存
    if os.path.exists("chunks.pkl"):
        os.remove("chunks.pkl")
    # 返回空状态：空对话、空输入、空引用、空Token、提示文字
    return [], "", "", "", "⚠️ 知识库已清空，请重新上传文档。", []


def clear_chat():
    """清空对话历史和所有显示"""
    return [], "", "", ""

# ========== 构建网页界面 ==========
with gr.Blocks(title="本地化知识库智能问答系统") as demo:
    gr.Markdown("# 📚 本地化知识库智能问答系统")
    gr.Markdown("上传文档后，即可基于文档内容进行多轮智能问答，支持 TXT / PDF / Word 格式，回答自带原文引用溯源。")
    
    with gr.Row():
        # 左侧：文档上传区
        with gr.Column(scale=1):
            gr.Markdown("### 第一步：上传文档")
            file_input = gr.File(label="选择文档", file_types=[".txt", ".pdf", ".docx"], file_count="multiple")
            upload_status = gr.Textbox(label="上传状态", interactive=False, lines=2)           
            gr.Markdown("### 消耗统计")
            token_output = gr.Textbox(label="", interactive=False, lines=6)
            clear_kb_btn = gr.Button("清空知识库", variant="stop")
            clear_btn = gr.Button("清空对话", variant="secondary")
        
        # 右侧：问答区
        with gr.Column(scale=2):
            gr.Markdown("### 第二步：对话提问")
            chatbot = gr.Chatbot(label="对话窗口", height=400)
            question_input = gr.Textbox(label="请输入你的问题", placeholder="例如：被告人叫什么名字？")
            
            with gr.Row():
                submit_btn = gr.Button("提交问题", variant="primary")
            
            gr.Markdown("### 引用来源")
            ref_output = gr.Textbox(label="", lines=8)
    
    # 上传文档事件
    file_input.upload(
        fn=build_vector_store,
        inputs=file_input,
        outputs=[upload_status, chatbot]
    )
    
    # 提交问题事件
    submit_btn.click(
        fn=answer_question,
        inputs=[question_input, chatbot],
        outputs=[chatbot, question_input, ref_output, token_output]
    )
    
    # 回车提交
    question_input.submit(
        fn=answer_question,
        inputs=[question_input, chatbot],
        outputs=[chatbot, question_input, ref_output, token_output]
    )
    
    # 新增：清空知识库事件
    clear_kb_btn.click(
    fn=clear_knowledge_base,
    inputs=[],
    outputs=[chatbot, question_input, ref_output, token_output, upload_status, file_input]
)

    # 清空对话事件
    clear_btn.click(
        fn=clear_chat,
        outputs=[chatbot, question_input, ref_output, token_output]
    )

# 启动网页
if __name__ == "__main__":
    demo.launch(inbrowser=True)
