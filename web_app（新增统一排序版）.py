import os
import shutil
from pathlib import Path
import gradio as gr
from numpy import dot
from numpy.linalg import norm

# ========== 配置区 ==========
os.environ["DASHSCOPE_API_KEY"] = "sk-ws-H.PERIMDR.VRKH.MEUCIBtgY2J-UZDsKe0aYEjI9hmDVb3nkpsb265cA-62kpyGAiEA1mH4hA3KA2cmr65n6DwmQQvZ4zst9YQDTvep5nqEvIA"
# 向量库临时保存路径
VECTOR_STORE_PATH = "temp_faiss_index"
# 对话历史保留轮数 & 查询改写使用轮数
MAX_HISTORY_ROUND = 2
# 语义相似度最低阈值（0-1，低于此值的段落直接过滤）
SIMILARITY_THRESHOLD = 0.4

# ========== 导入依赖 ==========
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import DashScopeEmbeddings
from pypdf import PdfReader
from rank_bm25 import BM25Okapi
from langchain_core.documents import Document
import docx
import dashscope
import pickle

# ========== 全局缓存 ==========
vector_store_cache = None
bm25_cache = None
all_chunks_cache = None

# ========== 工具函数 ==========
def load_document(file_path):
    """读取不同格式文档"""
    file_path = Path(file_path)
    suffix = file_path.suffix.lower()
    
    if suffix == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    
    elif suffix == ".pdf":
        text = ""
        reader = PdfReader(file_path)
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    
    elif suffix == ".docx":
        doc = docx.Document(file_path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        return text
    
    else:
        raise ValueError(f"不支持的文件格式：{suffix}")

def build_vector_store(file_obj):
    """根据上传的文档构建向量库，同时清空对话历史"""
    # 清理旧的向量库
    if os.path.exists(VECTOR_STORE_PATH):
        shutil.rmtree(VECTOR_STORE_PATH)
    if os.path.exists("chunks.pkl"):
        os.remove("chunks.pkl")
    
    # 初始化总分块列表
    all_chunks = []
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=50,
        separators=["\n\n", "\n", "。", "，", " "]
    )
    
    # 循环处理每个上传的文档
    for file_item in file_obj:
        raw_text = load_document(file_item.name)
        chunks = text_splitter.split_text(raw_text)
        all_chunks.extend(chunks)
    
    chunks = all_chunks
    
    # 构建向量库
    embeddings = DashScopeEmbeddings(model="text-embedding-v2")
    vector_store = FAISS.from_texts(chunks, embeddings)
    vector_store.save_local(VECTOR_STORE_PATH)
    
    # 保存文本块副本，供BM25关键词检索使用
    with open("chunks.pkl", "wb") as f:
        pickle.dump(chunks, f)
    
    # 更新全局缓存
    global vector_store_cache, bm25_cache, all_chunks_cache
    all_chunks_cache = chunks
    vector_store_cache = vector_store
    
    # 预构建BM25索引存入缓存
    tokenized_corpus = [list(doc) for doc in chunks]
    bm25_cache = BM25Okapi(tokenized_corpus)
    
    # 返回：上传状态 + 空对话历史
    return f"✅ 知识库构建完成！共处理 {len(file_obj)} 个文档，解析 {len(chunks)} 个文本块。", []

def answer_question(query, chat_history):
    # 空输入直接返回
    if not query.strip():
        yield chat_history, "", "", "", ""
        return
    
    # ========== 第1步：先显示用户问题 ==========
    chat_history.append({"role": "user", "content": query})
    yield chat_history, "", "", "", ""
    
    # ========== 第2步：检查知识库 ==========
    if not os.path.exists(VECTOR_STORE_PATH) or not os.path.exists("chunks.pkl"):
        chat_history.append({"role": "assistant", "content": "⚠️ 请先上传文档构建知识库！"})
        yield chat_history, "", "", "", ""
        return
    
    # ========== 第2.5步：查询改写（结合历史对话补全语义） ==========
    # 取最近几轮对话用于改写
    history_for_rewrite = chat_history[-MAX_HISTORY_ROUND * 2:]
    rewrite_history_text = ""
    for msg in history_for_rewrite:
        if msg["role"] == "user":
            rewrite_history_text += f"用户：{msg['content']}\n"
        else:
            rewrite_history_text += f"助手：{msg['content']}\n"
    
    # 改写提示词
    rewrite_prompt = f"""请结合对话历史，将用户当前的问题改写为一句完整、独立、无指代省略的检索问句。
要求：
1. 补全省略的主语、指代（如它、此案、这个人等），补充完整上下文信息
2. 只输出改写后的问题本身，不要任何解释、前缀、后缀
3. 保持原问题的核心语义，不要添加额外信息
对话历史：
{rewrite_history_text}
用户当前问题：{query}
改写后的问题："""
    
    # 调用大模型执行改写（非流式）
    rewrite_resp = dashscope.Generation.call(
        model="qwen-turbo",
        prompt=rewrite_prompt,
        result_format="text",
        stream=False
    )
    # 提取改写后的问题，去除首尾空白，空值兜底
    rewritten_query = rewrite_resp.output.text.strip()
    if not rewritten_query:
        rewritten_query = query
    print("【原问题】", query)
    print("【改写后】", rewritten_query)
    
    # ========== 第3步：混合检索上下文 ==========
    global vector_store_cache, bm25_cache, all_chunks_cache
    embeddings = DashScopeEmbeddings(model="text-embedding-v2")
    
    # 向量检索：优先读缓存
    if vector_store_cache is None:
        vector_store = FAISS.load_local(
            VECTOR_STORE_PATH, 
            embeddings,
            allow_dangerous_deserialization=True
        )
        vector_store_cache = vector_store
    else:
        vector_store = vector_store_cache
    vector_docs = vector_store.similarity_search(rewritten_query, k=3)
    
    # BM25检索：优先读缓存
    if bm25_cache is None or all_chunks_cache is None:
        with open("chunks.pkl", "rb") as f:
            all_chunks = pickle.load(f)
        all_chunks_cache = all_chunks
        tokenized_corpus = [list(doc) for doc in all_chunks]
        bm25_cache = BM25Okapi(tokenized_corpus)
        bm25 = bm25_cache  # 修复：给局部变量赋值
    else:
        all_chunks = all_chunks_cache
        bm25 = bm25_cache
    
    tokenized_query = list(rewritten_query)
    bm25_scores = bm25.get_scores(tokenized_query)
    top_bm25_idx = sorted(range(len(bm25_scores)), key=lambda i: bm25_scores[i], reverse=True)[:3]
    bm25_docs = [Document(page_content=all_chunks[i]) for i in top_bm25_idx]
    
    # 3. 合并去重
    combined_docs = []
    seen_contents = set()
    for doc in vector_docs + bm25_docs:
        if doc.page_content not in seen_contents:
            seen_contents.add(doc.page_content)
            combined_docs.append(doc)
    
    # 极端情况：无候选段落直接返回
    if not combined_docs:
        chat_history.append({"role": "assistant", "content": "知识库中暂无相关内容"})
        yield chat_history, "", "", "", ""
        return
    
    # ========== 重排序：统一语义相似度精排 ==========
    # 生成查询向量
    query_embedding = embeddings.embed_query(rewritten_query)
    
    # 计算每个候选段落和查询的余弦相似度
    scored_docs = []
    for doc in combined_docs:
        doc_embedding = embeddings.embed_query(doc.page_content)
        # 余弦相似度计算公式
        cos_sim = dot(query_embedding, doc_embedding) / (norm(query_embedding) * norm(doc_embedding))
        scored_docs.append((cos_sim, doc))
    
    # 按相似度降序排序
    scored_docs.sort(key=lambda x: x[0], reverse=True)
    
    # 过滤低于阈值的低质量段落
    filtered_docs = [(sim, doc) for sim, doc in scored_docs if float(sim) >= SIMILARITY_THRESHOLD]
    # 最多取Top4作为最终上下文
    related_docs = [doc for sim, doc in filtered_docs[:4]]
    
    # 兜底：过滤后无有效段落，直接返回无结果
    if not related_docs:
        chat_history.append({"role": "assistant", "content": "知识库中暂无相关内容"})
        yield chat_history, "", "", "", ""
        return
    
    # ========== 封装检索详情（Markdown美化格式） ==========
    # 1. 构建候选段落排序列表
    rank_content = ""
    for idx, (sim, doc) in enumerate(scored_docs, 1):
        # 截取前60字作为预览，去除换行
        preview_text = doc.page_content[:60].replace("\n", " ").strip() + "..."
        sim_score = round(float(sim) * 100, 1)  # 转成百分比更直观
        rank_content += f"#### 第 {idx} 名 · 语义匹配度 {sim_score}%\n"
        rank_content += f"> {preview_text}\n\n"
    
    # 2. 拼接完整的Markdown
    retrieve_detail_md = f"""
### 🔍 检索全流程详情
---
#### 1. 查询改写
| 类型 | 内容 |
| :--- | :--- |
| 原始问题 | {query} |
| 改写后问题 | {rewritten_query} |
---
#### 2. 候选段落语义排序（共 {len(scored_docs)} 段）
{rank_content}
---
✅ 最终筛选 Top **{len(related_docs)}** 段作为生成回答的上下文
"""
    
    # ========== 第4步：构建提示词 ==========
    # 给上下文段落编号
    context = ""
    for i, doc in enumerate(related_docs, 1):
        context += f"[{i}] {doc.page_content}\n\n"
    
    # 拼接历史对话
    history_text = ""
    history_messages = chat_history[-MAX_HISTORY_ROUND * 2:]
    for msg in history_messages:
        if msg["role"] == "user":
            history_text += f"用户：{msg['content']}\n"
        else:
            history_text += f"助手：{msg['content']}\n"
    
   # 最终提示词
    prompt = f"""请你结合对话历史，仅根据下面的上下文内容回答用户的问题。
规则：
1. 所有答案必须100%来自提供的上下文资料，禁止使用任何外部知识，禁止编造上下文没有的信息、人物、情节
2. 每个知识点回答后，必须用 [编号] 标注对应的来源段落，多个来源分别标注，如[1][2]
3. 如果上下文里完全没有相关信息，直接固定回答：知识库中暂无相关内容
4. 严格贴合问题角度回答，不额外发散，不重复无关内容
5. 结合对话历史理解用户的问题，保持回答的连贯性和一致性
对话历史：
{history_text}
上下文：
{context}
用户当前问题：{query}
请给出简洁准确的回答："""
    
    # ========== 第5步：流式调用大模型 ==========
    responses = dashscope.Generation.call(
        model="qwen-turbo",
        prompt=prompt,
        result_format="text",
        stream=True
    )
    answer = ""
    usage = None
    # 用副本操作，不污染原始对话历史
    temp_chat = chat_history.copy()
    # 预先占位一条空的助手消息
    temp_chat.append({"role": "assistant", "content": ""})
    
    for resp in responses:
        if resp.output is None:
            temp_chat[-1]["content"] = f"❌ 调用失败：{resp.message}"
            yield temp_chat, "", "", "", ""
            return
        
        # 直接使用最新全量文本，不要累加
        if resp.output.text:
            answer = resp.output.text
        
        # 记录Token统计
        if resp.usage:
            usage = resp.usage
        
        # 实时更新助手消息
        temp_chat[-1]["content"] = answer
        yield temp_chat, "", "", "", retrieve_detail_md
    
    # ========== 第6步：生成完成，返回完整结果 ==========
    # 同步最终对话
    chat_history[:] = temp_chat
    
    # 格式化引用原文
    ref_text = "📄 引用原文（与回答中编号对应）：\n\n"
    for i, doc in enumerate(related_docs, 1):
        ref_text += f"[{i}] {doc.page_content.strip()}\n\n"
    
    # 格式化Token统计
    token_text = f"""📊 Token消耗统计
输入Token：{usage.input_tokens}
输出Token：{usage.output_tokens}
本次总计：{usage.total_tokens} Tokens"""
    
    # 最终返回完整结果
    yield chat_history, "", ref_text, token_text, retrieve_detail_md

def clear_knowledge_base():
    """清空整个知识库"""
    global vector_store_cache, bm25_cache, all_chunks_cache
    
    if os.path.exists(VECTOR_STORE_PATH):
        shutil.rmtree(VECTOR_STORE_PATH)
    if os.path.exists("chunks.pkl"):
        os.remove("chunks.pkl")
    
    # 清空内存缓存
    vector_store_cache = None
    bm25_cache = None
    all_chunks_cache = None
    
    # 返回顺序对应前端outputs：chatbot, question_input, ref_output, token_output, upload_status, retrieve_detail_box, file_input
    return [], "", "", "", "⚠️ 知识库已清空，请重新上传文档。", "", None

def clear_chat():
    """清空对话历史和所有显示"""
    return [], "", "", "", ""

# ========== 构建网页界面 ==========
with gr.Blocks(title="本地化知识库智能问答系统") as demo:
    gr.Markdown("# 📚 本地化知识库智能问答系统")
    gr.Markdown("上传文档后，即可基于文档内容进行多轮智能问答，支持 TXT / PDF / Word 格式，回答自带原文引用溯源。")
    
    # 最外层：一行三列（所有列必须包裹在这同一个 Row 里面）
    with gr.Row():
        # ========== 第1列：左侧上传区 ==========
        with gr.Column(scale=1):
            gr.Markdown("### 第一步：上传文档")
            file_input = gr.File(label="选择文档", file_types=[".txt", ".pdf", ".docx"], file_count="multiple")
            upload_status = gr.Textbox(label="上传状态", interactive=False, lines=2)           
            gr.Markdown("### 消耗统计")
            token_output = gr.Textbox(label="", interactive=False, lines=6)
            clear_kb_btn = gr.Button("清空知识库", variant="stop")
            clear_btn = gr.Button("清空对话", variant="secondary")
        
        # ========== 第2列：中间问答区 ==========
        with gr.Column(scale=2):
            gr.Markdown("### 第二步：对话提问")
            chatbot = gr.Chatbot(label="对话窗口", height=400)
            question_input = gr.Textbox(label="请输入你的问题", placeholder="例如：被告人叫什么名字？")
            
            with gr.Row():
                submit_btn = gr.Button("提交问题", variant="primary")
            
            gr.Markdown("### 引用来源")
            ref_output = gr.Textbox(label="", lines=8)
        
        # ========== 第3列：右侧检索详情 ==========
        with gr.Column(scale=1):
            with gr.Accordion("🔍 检索过程详情", open=False):
                retrieve_detail_box = gr.Markdown(value="")
    
    # 上传文档事件
    file_input.upload(
        fn=build_vector_store,
        inputs=file_input,
        outputs=[upload_status, chatbot]
    )
    
    # 提交问题事件
    submit_btn.click(
        fn=answer_question,
        inputs=[question_input, chatbot],
        outputs=[chatbot, question_input, ref_output, token_output, retrieve_detail_box]
    )
    
    # 回车提交
    question_input.submit(
        fn=answer_question,
        inputs=[question_input, chatbot],
        outputs=[chatbot, question_input, ref_output, token_output, retrieve_detail_box]
    )
    
    # 清空知识库事件
    clear_kb_btn.click(
        fn=clear_knowledge_base,
        inputs=[],
        outputs=[chatbot, question_input, ref_output, token_output, upload_status, retrieve_detail_box, file_input]
    )
    
    # 清空对话事件
    clear_btn.click(
        fn=clear_chat,
        outputs=[chatbot, question_input, ref_output, token_output, retrieve_detail_box]
    )

# 启动网页
if __name__ == "__main__":
    demo.launch(inbrowser=True)
